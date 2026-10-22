"""Recupera entregas que estão no EPOD da Cainiao mas faltam no sistema.

Lê um ou mais ficheiros EPOD_TASK_LIST, encontra os waybills com uma linha
"Delivered" cuja (waybill, Task Date) não está como Delivered na BD, e
importa SÓ esses waybills com a lógica normal (import_operation_file).
Os outros waybills do ficheiro não são tocados.

Por omissão só mostra o que faria. Use --apply para gravar.

    python manage.py recover_epod_missing /tmp/a.xlsx /tmp/b.xlsx
    python manage.py recover_epod_missing /tmp/a.xlsx --apply
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime

from django.core.management.base import BaseCommand, CommandError
from openpyxl import Workbook


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


class Command(BaseCommand):
    help = "Importa só as entregas do EPOD que faltam na BD (dry-run por omissão)."

    def add_arguments(self, parser):
        parser.add_argument("files", nargs="+")
        parser.add_argument("--apply", action="store_true", help="Grava na BD.")

    def handle(self, *args, files, apply, **opts):
        from settlements.cainiao_views import (
            _find_header_row, _load_workbook_rows, import_operation_file,
        )
        from settlements.models import CainiaoOperationTask

        for path in files:
            with open(path, "rb") as fh:
                rows = _load_workbook_rows(fh.read())
            header, hidx = _find_header_row(rows, "Waybill Number")
            if header is None:
                raise CommandError(f"{path}: coluna 'Waybill Number' não encontrada.")
            h = [str(x or "").strip() for x in header]

            def col(*names):
                for n in names:
                    if n in h:
                        return h.index(n)
                return None

            iw, ist, itd = col("Waybill Number"), col("Task Status"), col("Task Date")
            icr = col("Driver", "Courier Name")
            if None in (iw, ist, itd):
                raise CommandError(f"{path}: faltam colunas Waybill/Task Status/Task Date.")

            by_wb = defaultdict(list)
            delivered = {}
            for r in rows[hidx + 1:]:
                wb = str(r[iw] or "").strip() if iw < len(r) else ""
                if not wb:
                    continue
                by_wb[wb].append(r)
                if str(r[ist] or "").strip() == "Delivered":
                    delivered[wb] = (_to_date(r[itd]), str(r[icr] or "").strip() if icr is not None else "")

            in_db = set(
                CainiaoOperationTask.objects.filter(
                    waybill_number__in=list(delivered), task_status="Delivered",
                ).values_list("waybill_number", "task_date")
            )
            missing = [wb for wb, (td, _c) in delivered.items() if td and (wb, td) not in in_db]

            # Ambíguo: "Cancel" com data posterior à entrega (pacote voltou ao
            # armazém?). Não se importa — fica para revisão manual.
            ambiguous = [
                wb for wb in missing
                if any(
                    str(r[ist] or "").strip() == "Cancel"
                    and (_to_date(r[itd]) or date.min) > delivered[wb][0]
                    for r in by_wb[wb]
                )
            ]
            missing = [wb for wb in missing if wb not in ambiguous]
            for wb in ambiguous:
                self.stdout.write(self.style.WARNING(
                    f"  revisão manual (Cancel depois da entrega): {wb} "
                    f"{delivered[wb][0]} {delivered[wb][1]}"
                ))

            per = Counter((delivered[wb][1], str(delivered[wb][0])) for wb in missing)
            self.stdout.write(f"{path}: {len(rows) - hidx - 1} linhas, {len(delivered)} Delivered, "
                              f"{len(missing)} em falta na BD")
            for (courier, d), n in sorted(per.items(), key=lambda x: (x[0][1], x[0][0])):
                self.stdout.write(f"  {d}  {courier:<35} {n}")

            if not missing or not apply:
                continue

            wbk = Workbook()
            ws = wbk.active
            ws.append(list(header))
            for wb in missing:
                for r in by_wb[wb]:
                    ws.append(list(r))
            buf = io.BytesIO()
            wbk.save(buf)
            payload, status = import_operation_file(
                buf.getvalue(), f"recover_{path.rsplit('/', 1)[-1]}", user=None,
            )
            if not payload.get("success"):
                raise CommandError(f"{path}: importação falhou ({status}): {payload.get('error')}")
            self.stdout.write(self.style.SUCCESS(
                f"  importado: {payload['total_novos']} novos, "
                f"{payload['total_atualizados']} actualizados, datas {payload['date_range']}"
            ))
