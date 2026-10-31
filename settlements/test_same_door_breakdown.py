import io
from datetime import date, datetime, timezone
from decimal import Decimal
from unittest import mock

from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import Workbook
from reportlab.platypus import Spacer, Table

from core.models import Partner
from drivers_app.models import DriverProfile

from .models import CainiaoOperationTask
from .services_driver_breakdown import driver_breakdown
from .services_same_door import normalize_address, same_door_extra_ids

LOCMEM = {"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}}
D1 = date(2026, 10, 2)


def _task(wb, addr="Rua A 10", zc="4740-013", stop="-1", status="Delivered",
          dtype="TO_DOOR", day=D1, courier="Rui_LF", hour=10, lat="41.1", lng="-8.1"):
    return CainiaoOperationTask.objects.create(
        waybill_number=wb, task_date=day, task_status=status, courier_name=courier,
        delivery_type=dtype, detailed_address=addr, zip_code=zc, task_group_order=stop,
        receiver_latitude=lat, receiver_longitude=lng,
        delivery_time=datetime(2026, 10, 2, hour, 0, tzinfo=timezone.utc),
    )


class NormalizeAddressTests(TestCase):
    def test_variantes_da_mesma_morada_ficam_iguais(self):
        self.assertEqual(normalize_address("Rua porto carreiro n 25"),
                         normalize_address("rua porto carreiro n 25 4740-013"))
        self.assertEqual(normalize_address("Rua Pereira Azurara, 58 58"),
                         normalize_address("Rua Pereira Azurara, 58"))
        self.assertEqual(normalize_address("Rua de cassufas nº670 Rua de cassufas nº670"),
                         normalize_address("Rua de Cassufas 670"))

    def test_andares_diferentes_nao_juntam(self):
        self.assertNotEqual(normalize_address("Rua X 10 1 esq"), normalize_address("Rua X 10 2 esq"))


@override_settings(CACHES=LOCMEM)
class SameDoorTests(TestCase):
    def setUp(self):
        self.p = Partner.objects.create(
            name="CAINIAO", nif="514154411", contact_email="c@x.pt",
            same_door_enabled=True, same_door_additional_price=Decimal("0.30"),
            pudo_enabled=True,
        )

    def test_morada_escrita_diferente_conta_como_mesma_porta(self):
        a = _task("W1", addr="Rua porto carreiro n 25", hour=9)
        b = _task("W2", addr="rua porto carreiro n 25 4740-013", hour=10)
        self.assertEqual(same_door_extra_ids([a.id, b.id], self.p), {b.id})

    def test_mesma_paragem_cainiao_junta_moradas_diferentes(self):
        a = _task("W1", addr="Av Poveiros 265 R/C", stop="12.0", hour=9)
        b = _task("W2", addr="Avenida Poveiros 265", stop="12.0", hour=10)
        c = _task("W3", addr="Outra rua 1", stop="13.0", hour=11)
        self.assertEqual(same_door_extra_ids([a.id, b.id, c.id], self.p), {b.id})

    def test_paragem_de_outro_motorista_nao_junta(self):
        a = _task("W1", addr="Rua 1", stop="5.0")
        b = _task("W2", addr="Rua 2", stop="5.0", courier="Ana_LF")
        self.assertEqual(same_door_extra_ids([a.id, b.id], self.p), set())

    def test_dias_diferentes_e_pudo_nao_contam(self):
        a = _task("W1")
        b = _task("W2", day=date(2026, 10, 3))
        c = _task("W3", dtype="PUDO")
        self.assertEqual(same_door_extra_ids([a.id, b.id, c.id], self.p), set())


@override_settings(CACHES=LOCMEM)
class DriverBreakdownTests(TestCase):
    def setUp(self):
        self.p = Partner.objects.create(
            name="CAINIAO", nif="514154411", contact_email="c@x.pt",
            same_door_enabled=True, same_door_additional_price=Decimal("0.30"),
            pudo_enabled=True,
        )
        self.d = DriverProfile.objects.create(
            nif="123456789", nome_completo="Rui", telefone="910000000",
            email="rui@x.pt", apelido="Rui_LF", price_per_package=Decimal("1.10"),
        )

    def test_total_primeiro_depois_divisao_e_valor(self):
        _task("W1", hour=9)
        _task("W2", hour=10)                       # mesma porta
        _task("W3", addr="Rua B 2")                # porta normal
        _task("W4", dtype="PUDO", lat="41.2")      # PUDO
        _task("W5", status="Attempt Failure", addr="Rua C 3")
        bd = driver_breakdown(self.d, D1, D1)
        t = bd["totals"]
        self.assertEqual((t["delivered"], t["door_normal"], t["same_door"], t["pudo"], t["failures"]),
                         (4, 2, 1, 1, 1))
        self.assertEqual(t["delivered"], t["door_normal"] + t["same_door"] + t["pudo"])
        # 2 × 1,10 + 1 × 0,30 + PUDO 1ª entrega (1,00)
        self.assertEqual(bd["amounts"]["total"], Decimal("2.20") + Decimal("0.30") + Decimal("1.00"))
        self.assertEqual(bd["days"][0]["delivered"], 4)
        # Horas da 1ª e última entrega do dia (ideia 1) e grupos da mesma porta (ideia 5)
        self.assertEqual((bd["days"][0]["first"].hour, bd["days"][0]["last"].hour), (10, 11))
        self.assertEqual([(g["n"], g["waybills"]) for g in bd["same_door_groups"]], [(2, ["W1", "W2"])])


@override_settings(CACHES=LOCMEM)
class RecoverEpodMissingTests(TestCase):
    def _xlsx(self, tmp_path_rows):
        wb = Workbook()
        ws = wb.active
        ws.append(["Waybill Number", "Task Date", "Task Status", "Driver"])
        for r in tmp_path_rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_dry_run_lista_so_o_que_falta_e_apply_importa_so_esses(self):
        CainiaoOperationTask.objects.create(
            waybill_number="OK1", task_date=date(2026, 9, 23), task_status="Delivered", courier_name="Rui_LF",
        )
        content = self._xlsx([
            ["OK1", "2026-09-23", "Delivered", "Rui_LF"],
            ["MISS1", "2026-09-23", "Delivered", "Rui_LF"],
            ["FAIL1", "2026-09-23", "Attempt Failure", "Rui_LF"],
        ])
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
            f.write(content)
            f.flush()
            out = io.StringIO()
            call_command("recover_epod_missing", f.name, stdout=out)
            self.assertIn("1 em falta", out.getvalue())
            with mock.patch("settlements.cainiao_views.import_operation_file",
                            return_value=({"success": True, "total_novos": 1, "total_atualizados": 0,
                                           "date_range": "2026-09-23"}, 200)) as imp:
                call_command("recover_epod_missing", f.name, "--apply", stdout=io.StringIO())
        sent = imp.call_args[0][0]
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(sent)).active.iter_rows(values_only=True))
        self.assertEqual([r[0] for r in rows[1:]], ["MISS1"])


@override_settings(CACHES=LOCMEM, SECURE_SSL_REDIRECT=False)
class BreakdownPagesTests(TestCase):
    """As 4 páginas mostram o total do motorista e a divisão."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from .models import DriverPreInvoice, PreInvoiceLine

        self.p = Partner.objects.create(
            name="CAINIAO", nif="514154411", contact_email="c@x.pt",
            same_door_enabled=True, same_door_additional_price=Decimal("0.30"), pudo_enabled=True,
        )
        self.d = DriverProfile.objects.create(
            nif="123456789", nome_completo="Rui", telefone="910000000",
            email="rui@x.pt", apelido="Rui_LF", price_per_package=Decimal("1.10"),
        )
        today = date.today()
        for i, (addr, dtype) in enumerate([("Rua A 1", "TO_DOOR"), ("Rua A 1", "TO_DOOR"), ("Rua B 2", "PUDO")]):
            CainiaoOperationTask.objects.create(
                waybill_number=f"W{i}", task_date=today, task_status="Delivered", courier_name="Rui_LF",
                delivery_type=dtype, detailed_address=addr, zip_code="4740-013",
            )
        self.pf = DriverPreInvoice.objects.create(
            driver=self.d, periodo_inicio=today, periodo_fim=today, status="CALCULADO",
        )
        PreInvoiceLine.objects.create(pre_invoice=self.pf, parceiro=self.p, total_pacotes=2,
                                      taxa_por_entrega=Decimal("1.10"))
        self.client.force_login(get_user_model().objects.create_superuser("adm", "a@x.pt", "x"))

    def test_paginas_mostram_total_e_divisao(self):
        from django.urls import reverse
        r = self.client.get(reverse("drivers_app:driver_portal", args=[self.d.id]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "mesma porta")
        r = self.client.get(reverse("drivers_app:driver_portal_invoices", args=[self.d.id]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Total entregue pelo motorista")
        self.assertContains(r, "Mesma porta")
        r = self.client.get(reverse("drivers_app:driver_pre_invoice_detail", args=[self.d.id, self.pf.id]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Entregas pagas")
        self.assertContains(r, "Mesma porta")
        self.assertContains(r, "Não confere com o sistema")  # PF tem 2, sistema tem 3
        r = self.client.get(reverse("core:partner-detail", args=[self.p.id]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Mesma porta")

    def test_paginas_tem_divs_equilibrados(self):
        """Um </div> a mais parte o layout (aconteceu no dashboard do portal)."""
        from html.parser import HTMLParser
        from django.urls import reverse

        class Count(HTMLParser):
            depth, lowest = 0, 0

            def handle_starttag(self, tag, attrs):
                if tag == "div":
                    self.depth += 1

            def handle_endtag(self, tag):
                if tag == "div":
                    self.depth -= 1
                    self.lowest = min(self.lowest, self.depth)

        for url in (
            reverse("drivers_app:driver_portal", args=[self.d.id]),
            reverse("drivers_app:driver_portal_invoices", args=[self.d.id]),
            reverse("drivers_app:driver_pre_invoice_detail", args=[self.d.id, self.pf.id]),
        ):
            c = Count()
            c.feed(self.client.get(url).content.decode())
            self.assertEqual((c.depth, c.lowest), (0, 0), url)


@override_settings(CACHES=LOCMEM)
class RecoverAmbiguousTests(TestCase):
    def test_cancel_depois_da_entrega_nao_e_importado(self):
        import tempfile
        wb = Workbook()
        ws = wb.active
        ws.append(["Waybill Number", "Task Date", "Task Status", "Driver"])
        ws.append(["AMB1", "2026-09-13", "Delivered", "Gus_LF"])
        ws.append(["AMB1", "2026-09-14", "Cancel", "Gus_LF"])
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
            wb.save(f.name)
            out = io.StringIO()
            with mock.patch("settlements.cainiao_views.import_operation_file") as imp:
                call_command("recover_epod_missing", f.name, "--apply", stdout=out)
        imp.assert_not_called()
        self.assertIn("revisão manual", out.getvalue())
        self.assertIn("0 em falta", out.getvalue())


@override_settings(CACHES=LOCMEM)
class PreInvoicePdfTests(TestCase):
    def test_pdf_tem_resumo_da_operacao_e_3_blocos(self):
        from .models import DriverPreInvoice, PreInvoiceLine
        from .reports import pf_pdf_sections
        from .reports.pdf_generator import PDFGenerator

        p = Partner.objects.create(name="CAINIAO", nif="514154411", contact_email="c@x.pt",
                                   same_door_enabled=True, same_door_additional_price=Decimal("0.30"))
        d = DriverProfile.objects.create(nif="123456789", nome_completo="Rui", telefone="910000000",
                                         email="rui@x.pt", apelido="Rui_LF", price_per_package=Decimal("1.10"))
        _task("W1", hour=9)
        _task("W2", hour=10)
        pf = DriverPreInvoice.objects.create(driver=d, periodo_inicio=D1, periodo_fim=D1)
        PreInvoiceLine.objects.create(pre_invoice=pf, parceiro=p, total_pacotes=2,
                                      taxa_por_entrega=Decimal("1.10"), base_entregas=Decimal("2.20"))
        pf.recalcular()
        with mock.patch.object(pf_pdf_sections, "operation_summary",
                               wraps=pf_pdf_sections.operation_summary) as op, \
             mock.patch.object(pf_pdf_sections, "financial_summary",
                               wraps=pf_pdf_sections.financial_summary) as fin:
            pdf = PDFGenerator().generate_pre_invoice_pdf(pf).getvalue()
        self.assertTrue(pdf.startswith(b"%PDF"))
        op.assert_called_once()
        self.assertEqual(op.call_args[0][2]["totals"]["delivered"], 2)
        fin.assert_called_once()
        self.assertEqual(fin.call_args.kwargs["n_delivered"], 2)
        self.assertIsNotNone(pf_pdf_sections.fallback_logo_path())


@override_settings(CACHES=LOCMEM, SECURE_SSL_REDIRECT=False)
class PortalMonthChartTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.d = DriverProfile.objects.create(nif="123456789", nome_completo="Rui", telefone="910000000",
                                              email="rui@x.pt", apelido="Rui_LF")
        self.client.force_login(get_user_model().objects.create_superuser("adm", "a@x.pt", "x"))

    def test_mes_corrente_por_omissao_e_escolha_de_mes(self):
        import calendar
        from django.urls import reverse
        today = date.today()
        url = reverse("drivers_app:driver_portal", args=[self.d.id])
        r = self.client.get(url)
        self.assertContains(r, "mês corrente")
        self.assertEqual(len(r.context["history_30d"]), calendar.monthrange(today.year, today.month)[1])
        self.assertEqual(r.context["history_30d"][0]["date"], today.replace(day=1))
        r = self.client.get(url + "?mes=2026-02")
        self.assertContains(r, "Fevereiro 2026")
        self.assertEqual(len(r.context["history_30d"]), 28)
        self.assertNotContains(r, "mês corrente</span>")
        r = self.client.get(url + "?chart_end=2026-09-15")  # links antigos
        self.assertContains(r, "Setembro 2026")

    def test_faturas_periodo_mes_corrente_e_dias_especificos(self):
        from datetime import timedelta
        from django.urls import reverse
        today = date.today()
        url = reverse("drivers_app:driver_portal_invoices", args=[self.d.id])
        r = self.client.get(url)
        self.assertEqual((r.context["period"], r.context["date_from"], r.context["date_to"]),
                         ("month", today.replace(day=1), today))
        r = self.client.get(url + "?period=prev_month")
        last = today.replace(day=1) - timedelta(days=1)
        self.assertEqual((r.context["date_from"], r.context["date_to"]), (last.replace(day=1), last))
        r = self.client.get(url + "?period=month&mes=2026-02")
        self.assertEqual((r.context["date_from"], r.context["date_to"]), (date(2026, 2, 1), date(2026, 2, 28)))
        r = self.client.get(url + "?date_from=2026-09-21&date_to=2026-09-27")
        self.assertEqual((r.context["period"], r.context["date_from"], r.context["date_to"]),
                         ("custom", date(2026, 9, 21), date(2026, 9, 27)))
        self.assertContains(r, "21/09/2026 → 27/09/2026")


@override_settings(CACHES=LOCMEM)
class PdfIdeasTests(TestCase):
    """Ideias 1, 2, 4 e 5 do PDF: €/hora, falhas em €, PF anterior, mesma porta."""

    def setUp(self):
        from .models import DriverPreInvoice, PreInvoiceLine
        self.p = Partner.objects.create(name="CAINIAO", nif="514154411", contact_email="c@x.pt",
                                        same_door_enabled=True, same_door_additional_price=Decimal("0.30"))
        self.d = DriverProfile.objects.create(nif="123456789", nome_completo="Rui", telefone="910000000",
                                              email="rui@x.pt", apelido="Rui_LF", price_per_package=Decimal("1.10"))
        _task("W1", hour=9)
        _task("W2", hour=10)
        _task("W3", addr="Rua B 2", hour=12, status="Attempt Failure")
        _task("OLD", day=date(2026, 9, 25), addr="Rua Z 9")
        self.prev = DriverPreInvoice.objects.create(numero="PF-T1", driver=self.d, periodo_inicio=date(2026, 9, 21),
                                                    periodo_fim=date(2026, 9, 27), status="PAGO")
        PreInvoiceLine.objects.create(pre_invoice=self.prev, parceiro=self.p, total_pacotes=1,
                                      taxa_por_entrega=Decimal("1.10"), base_entregas=Decimal("1.10"))
        self.prev.recalcular()
        self.pf = DriverPreInvoice.objects.create(numero="PF-T2", driver=self.d, periodo_inicio=D1, periodo_fim=D1)
        PreInvoiceLine.objects.create(pre_invoice=self.pf, parceiro=self.p, total_pacotes=1,
                                      taxa_por_entrega=Decimal("1.20"), base_entregas=Decimal("1.20"))
        PreInvoiceLine.objects.create(pre_invoice=self.pf, parceiro=self.p, total_pacotes=1,
                                      taxa_por_entrega=Decimal("0.30"), base_entregas=Decimal("0.30"),
                                      observacoes="Preço especial: Pacote adicional na mesma porta")
        self.pf.recalcular()

    def test_horas_falhas_anterior_e_mesma_porta(self):
        from .reports import pf_pdf_sections as s
        from .services_driver_breakdown import driver_breakdown

        bd = driver_breakdown(self.d, D1, D1)
        day = bd["days"][0]
        self.assertEqual(s.hours_worked(day), Decimal(1))  # 09:00 → 10:00
        self.assertEqual(s.pf_rates(self.pf), (Decimal("1.20"), Decimal("0.30")))
        self.assertEqual(bd["totals"]["failures"], 1)  # 1 × 1,20 não ganhos
        comp = s.previous_comparison(self.pf, bd, s.day_values(self.pf, bd["days"]))
        self.assertTrue(comp)
        sd = s.same_door_section(self.pf, lambda t, c=None: Spacer(1, 1), bd)
        tbl = next(x for x in sd if isinstance(x, Table))
        self.assertEqual(tbl._cellvalues[1][2], "1 + 1")
        self.assertEqual(tbl._cellvalues[1][3], "€ 1,50")  # 1,20 + 0,30

    def test_pdf_completo_gera(self):
        from .reports.pdf_generator import PDFGenerator
        self.assertTrue(PDFGenerator().generate_pre_invoice_pdf(self.pf).getvalue().startswith(b"%PDF"))
